from __future__ import annotations

import io
import re
from typing import Iterable

from app.core.config import get_settings

SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5

RED = (208, 24, 32)
DARK = (26, 32, 44)
MID_GREY = (189, 189, 189)
LIGHT_GREY = (242, 242, 242)
BORDER_GREY = (220, 220, 220)
TEXT_GREY = (94, 94, 94)
MINT = (108, 229, 168)
WHITE = (255, 255, 255)
BLACK = (0, 0, 0)
SOFT_PINK = (244, 208, 214)
SOFT_GREEN = (39, 203, 101)
SOFT_GREY = (111, 123, 133)
SOFT_BLUE = (196, 240, 246)
SOFT_TAUPE = (196, 180, 160)
LIGHT_GREEN = (196, 249, 207)
LIGHT_PANEL = (239, 239, 239)
LIGHT_BLUE = (227, 247, 252)
LIGHT_TAUPE = (238, 233, 227)

OBJECTIVE_COLORS = [
    (29, 201, 95),
    (95, 95, 95),
    (193, 240, 240),
    (214, 203, 193),
]


Presentation = None
RGBColor = None
MSO_AUTO_SHAPE_TYPE = None
PP_ALIGN = None
MSO_ANCHOR = None
Inches = None
Pt = None


def _ensure_pptx() -> None:
    global Presentation, RGBColor, MSO_AUTO_SHAPE_TYPE, PP_ALIGN, MSO_ANCHOR, Inches, Pt
    if Presentation is not None:
        return

    try:
        from pptx import Presentation as _Presentation
        from pptx.dml.color import RGBColor as _RGBColor
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as _MSO_AUTO_SHAPE_TYPE
        from pptx.enum.text import PP_ALIGN as _PP_ALIGN, MSO_ANCHOR as _MSO_ANCHOR
        from pptx.util import Inches as _Inches, Pt as _Pt
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "PowerPoint generation requires 'python-pptx'. Install dependencies from requirements.txt."
        ) from exc

    Presentation = _Presentation
    RGBColor = _RGBColor
    MSO_AUTO_SHAPE_TYPE = _MSO_AUTO_SHAPE_TYPE
    PP_ALIGN = _PP_ALIGN
    MSO_ANCHOR = _MSO_ANCHOR
    Inches = _Inches
    Pt = _Pt


def _rgb(color):
    _ensure_pptx()
    return RGBColor(*color)


def _add_box(slide, left, top, width, height, fill, line=None, rounded=False):
    _ensure_pptx()
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if rounded
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(
        shape_type, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line or fill)
    return shape


def _add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=10,
    bold=False,
    color=BLACK,
    align=None,
    italic=False,
    margin=0.04,
    valign=None,
):
    _ensure_pptx()
    align = align or PP_ALIGN.LEFT
    valign = valign or MSO_ANCHOR.TOP
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = Inches(margin)
    box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin)
    box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    run.font.name = get_settings().pptx_font
    return box


def _add_bullets(
    slide,
    items: Iterable[str],
    left,
    top,
    width,
    height,
    *,
    size=8.5,
    color=TEXT_GREY,
    bullet="•",
):
    _ensure_pptx()
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)

    cleaned = [item.strip() for item in items if item and item.strip()]
    if not cleaned:
        cleaned = ["-"]

    first = True
    for item in cleaned:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = f"{bullet} {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(color)
        run.font.name = get_settings().pptx_font

    return box


def _text_list(items: Iterable[str], fallback: str = "-") -> list[str]:
    values = [item.strip() for item in items if item and item.strip()]
    return values or [fallback]


def _shorten(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1].rstrip() + "…"


def _font_size_for_length(text: str, base: float, thresholds: list[tuple[int, float]]) -> float:
    for max_len, size in thresholds:
        if len(text) <= max_len:
            return size
    return base


def _wrap_metric_value(label: str, value: str) -> str:
    value = str(value).strip()
    if label == "Budget":
        if " - " in value:
            return value.replace(" - ", "\n", 1)
        if value.count("-") == 1 and len(value) > 12:
            left, right = value.split("-", 1)
            return f"{left.strip()}-\n{right.strip()}"
    if len(value) > 16 and " / " in value:
        return value.replace(" / ", "\n", 1)
    return value


def _parse_duration_to_quarters(duration: str) -> int:
    months = re.search(r"(\d+)\s*months?", duration, re.IGNORECASE)
    if months:
        return max(1, round(int(months.group(1)) / 3))

    days = re.search(r"(\d+)\s*days?", duration, re.IGNORECASE)
    if days:
        return 1 if int(days.group(1)) <= 90 else 2

    return 1


def _parse_quarter_label(timeframe: str) -> tuple[int, int | None]:
    match = re.search(r"Q([1-4])(?:\s*(20\d{2}))?", timeframe, re.IGNORECASE)
    if not match:
        return 1, None
    quarter = int(match.group(1))
    year = int(match.group(2)) if match.group(2) else None
    return quarter, year


def _infer_security_domain(title: str, themes: list[str]) -> str:
    text = f"{title} {' '.join(themes)}".lower()
    if "third-party" in text or "vendor" in text:
        return "Third-Party Risk"
    if "iam" in text or "access" in text or "identity" in text:
        return "Identity and Access Management"
    if "encrypt" in text or "privacy" in text or "data" in text:
        return "Data Protection"
    if "cloud" in text:
        return "Cloud Security"
    if themes:
        return themes[0].split(":")[0][:36]
    return "Cybersecurity"


def _initiative_code(pillar_label: str, item_index: int) -> str:
    initial = next((char.upper() for char in pillar_label if char.isalpha()), "I")
    return f"{initial}{item_index + 1}"


def _clean_label(text: str, max_chars: int = 40) -> str:
    text = str(text or "").strip()
    if ":" in text:
        text = text.split(":", 1)[0].strip()
    return _shorten(text, max_chars)


def _pillarize_text(text: str, max_chars: int = 34) -> str:
    source = str(text or "").strip()
    lowered = source.lower()

    keyword_rules = [
        (("governance", "grc", "compliance"), "Agile Governance and Strategy"),
        (("customer", "trust", "digital"), "Enable the Business Products"),
        (("core", "platform", "modernize", "operations"), "Build a Strong Core Platforms"),
        (("financial", "data"), "Protect Financial Data"),
        (("access", "identity", "iam", "authentication"), "Strengthen Access Security"),
        (("cloud", "configuration"), "Secure Cloud Operations"),
        (("resilience", "recovery", "continuity"), "Elevate Resilience and Recovery"),
        (("vendor", "third-party", "partner"), "Strengthen Third-Party Governance"),
        (("privacy", "pii", "gdpr"), "Protect Privacy and Compliance"),
        (("data", "encrypt"), "Protect Data and Encryption"),
    ]
    for keywords, label in keyword_rules:
        if any(keyword in lowered for keyword in keywords):
            return _shorten(label, max_chars)

    source = source.rstrip(".")
    source = re.sub(r"^(ensure|build|enable|protect|improve|modernize|strengthen|safeguard|elevate)\s+", "", source, flags=re.IGNORECASE)
    words = source.split()
    if len(words) > 5:
        source = " ".join(words[:5])
    return _shorten(source, max_chars)


def _unique_preserve_order(items: Iterable[str]) -> list[str]:
    seen = set()
    ordered = []
    for item in items:
        value = str(item or "").strip()
        key = value.lower()
        if value and key not in seen:
            seen.add(key)
            ordered.append(value)
    return ordered


def _keyword_score(text: str, label: str) -> int:
    text_tokens = {
        token
        for token in re.findall(r"[a-zA-Z0-9]+", text.lower())
        if len(token) > 3
    }
    label_tokens = {
        token
        for token in re.findall(r"[a-zA-Z0-9]+", label.lower())
        if len(token) > 3
    }
    return len(text_tokens & label_tokens)


def _group_initiatives_by_pillars(
    initiative_details: list[dict], pillar_labels: list[str]
) -> list[list[str]]:
    grouped = [[] for _ in pillar_labels]
    if not pillar_labels:
        return grouped
    max_per_bucket = max(1, (len(initiative_details) + len(pillar_labels) - 1) // len(pillar_labels))

    for idx, item in enumerate(initiative_details):
        text = f"{item.get('title', '')} {item.get('rationale', '')}"
        scores = [_keyword_score(text, pillar) for pillar in pillar_labels]
        best_score = max(scores) if scores else 0
        if best_score > 0 and scores.count(best_score) == 1:
            best_index = scores.index(best_score)
        else:
            best_index = idx % len(pillar_labels)
        if len(grouped[best_index]) >= max_per_bucket:
            best_index = min(range(len(grouped)), key=lambda bucket_idx: len(grouped[bucket_idx]))
        grouped[best_index].append(item.get("title", "Initiative"))

    return grouped


def _priority_rank(priority: str) -> int:
    return {"High": 0, "Medium": 1, "Low": 2}.get(str(priority or "").title(), 3)


def _group_initiatives_for_reference_layout(
    initiative_details: list[dict],
) -> list[list[str]]:
    ordered = sorted(
        initiative_details,
        key=lambda item: (_priority_rank(item.get("priority", "")), item.get("title", "")),
    )
    titles = [item.get("title", "Initiative") for item in ordered]
    capacities = [2, 5, 4, 2]
    groups = [[] for _ in capacities]
    for idx, title in enumerate(titles):
        preferred = idx % len(groups)
        if len(groups[preferred]) < capacities[preferred]:
            groups[preferred].append(title)
            continue
        for bucket_idx, capacity in enumerate(capacities):
            if len(groups[bucket_idx]) < capacity:
                groups[bucket_idx].append(title)
                break
    return groups


def build_strategy_overview_slide_data_from_result(result: dict) -> dict:
    initiative_details = result.get("initiative_details", [])
    technology_pillars = [
        "Build a Strong Core Platforms",
        "Enable the Business Products",
        "Become a Global Benchmark DH 2.X",
    ]
    cybersecurity_pillars = [
        "Agile Governance and Strategy",
        "Modernize Security Operations and Technology",
        "Elevate Resilience and Risk Management",
        "Nurture People and Partnerships",
    ]
    initiative_groups = _group_initiatives_for_reference_layout(initiative_details)

    return {
        "technology_pillars": technology_pillars[:3],
        "cybersecurity_pillars": cybersecurity_pillars[:4],
        "initiative_groups": initiative_groups[:4],
    }




def build_initiative_slide_request_from_result(
    result: dict, request: dict | None = None
) -> dict:
    request = request or {}
    initiative_details = result.get("initiative_details", [])
    initiative_index = request.get("initiative_index", 0)
    if initiative_index < 0 or initiative_index >= len(initiative_details):
        raise IndexError("initiative_index is out of range for the selected result.")

    detail = initiative_details[initiative_index]
    title = detail.get("title", "Cybersecurity Initiative")
    recurring_themes = result.get("recurring_themes", [])
    business_alignments = result.get("business_alignments", [])
    board_summary = result.get("board_summary", {})

    matching_alignments = [
        item
        for item in business_alignments
        if item.get("initiative", "").strip().lower() == title.strip().lower()
    ]
    objectives = [
        item.get("business_objective", "")
        for item in matching_alignments
        if item.get("business_objective")
    ]
    dependencies = [
        item.get("dependency", "")
        for item in matching_alignments
        if item.get("dependency")
    ]

    overview = request.get("initiative_overview") or detail.get("rationale") or (
        board_summary.get("executive_summary", "")
    )
    timeframe = detail.get("timeframe", "")
    start_label = request.get("start_date") or timeframe.split("/")[0].strip() or "TBD"
    duration = request.get("duration") or (
        timeframe.split("/")[1].strip() if "/" in timeframe else timeframe or "TBD"
    )

    roadmap_items = []
    for idx, entry in enumerate(board_summary.get("roadmap", [])[:3]):
        quarter, year = _parse_quarter_label(entry)
        label = entry.split(":", 1)[1].strip() if ":" in entry else entry
        start_quarter = quarter if year is None else ((year - (year or year)) * 4) + quarter
        if year is not None:
            base_year = _parse_quarter_label(board_summary.get("roadmap", [entry])[0])[1] or year
            start_quarter = ((year - base_year) * 4) + quarter
        roadmap_items.append(
            {
                "name": _shorten(label, 32),
                "start_quarter": max(1, min(12, start_quarter)),
                "end_quarter": max(
                    1,
                    min(
                        12,
                        start_quarter
                        + (
                            _parse_duration_to_quarters(duration)
                            if idx == 0
                            else 1
                        )
                        - 1,
                    ),
                ),
            }
        )

    if not roadmap_items:
        roadmap_items = [
            {
                "name": _shorten(title, 32),
                "start_quarter": 1,
                "end_quarter": min(12, _parse_duration_to_quarters(duration)),
            }
        ]

    default_request = {
        "security_domain": _infer_security_domain(title, recurring_themes),
        "initiative_name": title,
        "start_date": start_label,
        "duration": duration,
        "budget": detail.get("cost_estimate", "TBD"),
        "ftes": request.get("ftes") or "TBD",
        "strategic_objectives": objectives[:4]
        or [opt.split(":", 1)[0] for opt in result.get("strategy_options", [])[:3]],
        "initiative_overview": overview,
        "initiative_owner": detail.get("owner", "CISO"),
        "supporting_functions": request.get("supporting_functions")
        or [detail.get("owner", "CISO"), "IT", "Risk", "Legal"],
        "initiative_projects": request.get("initiative_projects")
        or [{"name": title, "description": detail.get("rationale", "")}],
        "timeline_projects": request.get("timeline_projects") or roadmap_items,
        "ai_use_case": request.get("ai_use_case")
        or "Use AI to synthesize evidence, map risks, and prepare board-ready updates for this initiative.",
        "priority": detail.get("priority", "Medium"),
        "dependencies": request.get("dependencies") or dependencies,
        "challenges": request.get("challenges")
        or [_shorten(item, 80) for item in recurring_themes[:2]],
        "domains_impacted": request.get("domains_impacted")
        or [_shorten(item.split(":")[0], 28) for item in recurring_themes[:2]],
        "maturity_uplift": request.get("maturity_uplift")
        or ["x1.2", "x1.5", "x1.8"],
        "file_name": request.get("file_name") or f"{title.lower().replace(' ', '_')}.pptx",
    }

    for key, value in request.items():
        if key != "initiative_index" and value not in (None, "", []):
            default_request[key] = value

    return default_request


def generate_initiative_slide_ppt(slide_data: dict) -> bytes:
    _ensure_pptx()
    prs = _new_presentation()
    _render_initiative_slide(prs, slide_data)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def generate_initiative_slides_ppt_from_result(result: dict) -> tuple[bytes, str]:
    initiative_details = result.get("initiative_details", [])
    if not initiative_details:
        raise ValueError("This strategy result has no initiative_details to build slides from.")

    _ensure_pptx()
    prs = _new_presentation()
    _render_strategy_overview_slide(prs, build_strategy_overview_slide_data_from_result(result))
    for index in range(len(initiative_details)):
        slide_data = build_initiative_slide_request_from_result(
            result, {"initiative_index": index}
        )
        _render_initiative_slide(prs, slide_data)

    buffer = io.BytesIO()
    prs.save(buffer)

    analysis_name = result.get("analysis_name", "strategy_result")
    safe_name = re.sub(r"[^a-zA-Z0-9_-]+", "_", analysis_name).strip("_") or "strategy_result"
    return buffer.getvalue(), f"{safe_name}_initiative_slides.pptx"


def _new_presentation():
    _ensure_pptx()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH)
    prs.slide_height = Inches(SLIDE_HEIGHT)
    return prs


def _render_initiative_slide(prs, slide_data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _add_box(slide, 0.08, 0.2, 13.15, 6.9, WHITE, BORDER_GREY)

    _draw_left_panel(slide, slide_data)
    _draw_center_panel(slide, slide_data)
    _draw_right_panel(slide, slide_data)

    _add_text(
        slide,
        10.85,
        7.02,
        2.0,
        0.18,
        get_settings().pptx_copyright,
        size=5.5,
        color=TEXT_GREY,
        align=PP_ALIGN.RIGHT,
    )


def _render_strategy_overview_slide(prs, slide_data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_box(slide, 0.08, 0.2, 13.15, 6.9, WHITE, BORDER_GREY)
    _add_text(slide, 0.14, 0.28, 1.5, 0.1, "Classification: Confidential", size=6.2, color=BLACK)
    _add_text(
        slide,
        0.36,
        0.55,
        7.5,
        0.28,
        "Enabling Organizational Resilience Through Integrated Strategic Pillars",
        size=13,
        bold=True,
        color=DARK,
    )
    _draw_overview_content(slide, slide_data)
    _add_text(
        slide,
        10.85,
        7.02,
        2.0,
        0.18,
        get_settings().pptx_copyright,
        size=5.5,
        color=TEXT_GREY,
        align=PP_ALIGN.RIGHT,
    )


def _draw_overview_content(slide, data: dict) -> None:
    _ensure_pptx()

    COL_RING_CX = 1.53
    COL_RING_CY = 4.41
    RING_OUTER = 1.82
    RING_INNER_MARGIN = 0.32

    COL_TECH = 3.05
    TECH_W = 1.58
    TECH_H = 0.52

    COL_VLINE = 4.95
    COL_PILLAR = 5.18
    PILLAR_W = 1.88

    COL_I1 = 7.22
    COL_I2 = 9.18
    INIT_W = 1.88
    INIT_H = 0.36

    HDR_Y = 0.98
    HDR_H = 0.56

    ROW_CY = [2.3, 3.9, 5.5, 6.75]

    HEADER_GREY = (82, 82, 82)
    PILLAR_COLORS = [SOFT_GREEN, SOFT_GREY, SOFT_BLUE, SOFT_TAUPE]
    INIT_FILLS = [LIGHT_GREEN, LIGHT_PANEL, LIGHT_BLUE, LIGHT_TAUPE]

    headers = [
        (COL_TECH, HDR_Y, TECH_W, "Technology\nPillars"),
        (COL_PILLAR, HDR_Y, PILLAR_W, "Cybersecurity\nPillars"),
        (COL_I1, HDR_Y, INIT_W + (COL_I2 - COL_I1) + INIT_W, "Cybersecurity Initiatives"),
    ]
    for left, top, width, text in headers:
        _add_box(slide, left, top, width, HDR_H, HEADER_GREY, HEADER_GREY, rounded=True)
        _add_text(
            slide, left, top + 0.08, width, 0.35, text,
            size=9 if width < 3 else 10, bold=True, color=WHITE,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
        )
    sep1_x = COL_TECH + TECH_W + (COL_PILLAR - COL_TECH - TECH_W) / 2 - 0.06
    sep2_x = COL_PILLAR + PILLAR_W + (COL_I1 - COL_PILLAR - PILLAR_W) / 2 - 0.06
    for x in [sep1_x, sep2_x]:
        _add_text(slide, x, HDR_Y + 0.14, 0.12, 0.2, ">", size=14, bold=True, color=MID_GREY, align=PP_ALIGN.CENTER)

    ring_left = COL_RING_CX - RING_OUTER / 2
    ring_top = COL_RING_CY - RING_OUTER / 2
    outer = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(ring_left), Inches(ring_top), Inches(RING_OUTER), Inches(RING_OUTER),
    )
    outer.fill.solid()
    outer.fill.fore_color.rgb = _rgb(RED)
    outer.line.color.rgb = _rgb(RED)

    inner_d = RING_OUTER - RING_INNER_MARGIN * 2
    inner = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(ring_left + RING_INNER_MARGIN), Inches(ring_top + RING_INNER_MARGIN),
        Inches(inner_d), Inches(inner_d),
    )
    inner.fill.solid()
    inner.fill.fore_color.rgb = _rgb(WHITE)
    inner.line.color.rgb = _rgb(WHITE)

    cut_left = COL_RING_CX + inner_d * 0.18
    cut_top = COL_RING_CY - 0.38
    cut_w = RING_OUTER / 2
    cut_h = 0.76
    cutout = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(cut_left), Inches(cut_top), Inches(cut_w), Inches(cut_h),
    )
    cutout.fill.solid()
    cutout.fill.fore_color.rgb = _rgb(WHITE)
    cutout.line.color.rgb = _rgb(WHITE)

    _add_text(
        slide, COL_RING_CX - 0.55, COL_RING_CY - 0.22, 1.1, 0.44,
        "Cybersecurity\nStrategy", size=9, bold=True, color=DARK,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.0,
    )

    line_left = cut_left + 0.06
    for dy in [-0.16, 0.0, 0.16]:
        _add_box(slide, line_left, COL_RING_CY + dy - 0.013, 0.32, 0.026, RED, RED)

    slide.shapes.add_connector(
        1, Inches(COL_TECH - 0.3), Inches(1.82), Inches(COL_TECH - 0.3), Inches(ROW_CY[0] - 0.1)
    ).line.color.rgb = _rgb(BORDER_GREY)
    slide.shapes.add_connector(
        1, Inches(COL_TECH - 0.3), Inches(1.82), Inches(COL_TECH), Inches(1.82)
    ).line.color.rgb = _rgb(BORDER_GREY)

    technology_pillars = data.get("technology_pillars", [])[:3]
    tech_rows = [ROW_CY[0], ROW_CY[1], ROW_CY[3]]
    for idx in range(3):
        cy = tech_rows[idx]
        top = cy - TECH_H / 2
        _add_box(slide, COL_TECH, top, TECH_W, TECH_H, SOFT_PINK, SOFT_PINK)
        _add_box(slide, COL_TECH + TECH_W - 0.04, top, 0.04, TECH_H, RED, RED)
        label = technology_pillars[idx] if idx < len(technology_pillars) else f"Pillar {idx + 1}"
        _add_text(
            slide, COL_TECH + 0.04, top + 0.04, TECH_W - 0.12, TECH_H - 0.08,
            _shorten(label, 38),
            size=_font_size_for_length(label, 7.0, [(18, 8.4), (28, 7.6), (38, 6.8)]),
            bold=True, color=DARK, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.01,
        )

    vline_top = ROW_CY[0] - 0.22
    vline_bot = ROW_CY[3] + 0.22
    slide.shapes.add_connector(
        1, Inches(COL_VLINE), Inches(vline_top), Inches(COL_VLINE), Inches(vline_bot)
    ).line.color.rgb = _rgb(BORDER_GREY)

    cybersecurity_pillars = data.get("cybersecurity_pillars", [])[:4]
    pillar_heights = [0.48, 0.92, 0.92, 0.48]
    for idx in range(4):
        cy = ROW_CY[idx]
        ph = pillar_heights[idx]
        top = cy - ph / 2
        color = PILLAR_COLORS[idx]

        marker = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            Inches(COL_VLINE - 0.09), Inches(cy - 0.09), Inches(0.18), Inches(0.18),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = _rgb(WHITE)
        marker.line.color.rgb = _rgb(color)
        marker.line.width = Pt(2)

        _add_box(slide, COL_PILLAR, top, PILLAR_W, ph, color, color, rounded=True)
        label = cybersecurity_pillars[idx] if idx < len(cybersecurity_pillars) else f"Pillar {idx + 1}"
        _add_text(
            slide, COL_PILLAR + 0.08, top + 0.06, PILLAR_W - 0.16, ph - 0.12,
            _shorten(label, 40),
            size=_font_size_for_length(label, 7.2, [(20, 9), (30, 8.2), (40, 7.1)]),
            bold=True, color=WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.01,
        )

    groups = data.get("initiative_groups", [])[:4]
    pillars_for_code = data.get("cybersecurity_pillars", [])[:4]
    max_cols = 2
    row_capacities = [2, 5, 4, 2]
    gap_v = 0.04
    for row_idx in range(4):
        items = groups[row_idx] if row_idx < len(groups) else []
        fill = INIT_FILLS[row_idx]
        pillar_label = pillars_for_code[row_idx] if row_idx < len(pillars_for_code) else "I"
        n_items = min(len(items), row_capacities[row_idx])
        if n_items == 0:
            continue
        n_rows_actual = (n_items + max_cols - 1) // max_cols
        block_h = n_rows_actual * INIT_H + (n_rows_actual - 1) * gap_v
        block_top = ROW_CY[row_idx] - block_h / 2

        for s_idx in range(n_items):
            s_row = s_idx // max_cols
            s_col = s_idx % max_cols
            s_left = COL_I1 if s_col == 0 else COL_I2
            s_top = block_top + s_row * (INIT_H + gap_v)

            title = _shorten(items[s_idx], 44)
            code = _initiative_code(pillar_label, s_idx)
            _add_box(slide, s_left, s_top, INIT_W, INIT_H, fill, fill, rounded=True)
            _add_text(
                slide, s_left + 0.04, s_top + 0.03, INIT_W - 0.08, INIT_H - 0.06,
                f"{code}. {title}",
                size=_font_size_for_length(title, 5.9, [(18, 6.4), (28, 5.9), (44, 5.3)]),
                bold=True, color=DARK, valign=MSO_ANCHOR.MIDDLE, margin=0.01,
            )


def _draw_left_panel(slide, data: dict) -> None:
    _add_box(slide, 0.18, 0.3, 2.45, 0.82, RED)
    _add_text(slide, 0.35, 0.38, 1.7, 0.16, "Security Domain", size=7, color=WHITE)
    _add_text(
        slide,
        0.35,
        0.6,
        2.0,
        0.24,
        _shorten(data.get("security_domain", "Security Domain"), 34),
        size=13,
        bold=True,
        color=WHITE,
    )

    _add_box(slide, 0.18, 1.16, 2.45, 0.86, MID_GREY)
    _add_text(slide, 0.35, 1.3, 1.8, 0.16, "Initiative Name", size=7, color=WHITE)
    initiative_name = str(data.get("initiative_name", "Initiative Name")).strip()
    _add_text(
        slide,
        0.35,
        1.46,
        2.12,
        0.42,
        _shorten(initiative_name, 70),
        size=_font_size_for_length(
            initiative_name,
            8.0,
            [(24, 12), (34, 10.5), (50, 9.2), (70, 8.4)],
        ),
        bold=True,
        color=WHITE,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0.02,
    )

    metric_y = 2.15
    metrics = [
        ("Start Date", data.get("start_date", "TBD")),
        ("Duration", data.get("duration", "TBD")),
        ("Budget", data.get("budget", "TBD")),
        ("FTEs", data.get("ftes", "TBD")),
    ]
    for idx, (label, value) in enumerate(metrics):
        col = idx % 2
        row = idx // 2
        left = 0.18 + col * 1.225
        top = metric_y + row * 0.53
        _add_box(slide, left, top, 1.19, 0.5, WHITE, BORDER_GREY)
        _add_text(slide, left + 0.06, top + 0.06, 1.0, 0.12, label, size=7, color=TEXT_GREY)
        metric_value = _wrap_metric_value(label, str(value))
        _add_text(
            slide,
            left + 0.06,
            top + 0.19,
            1.05,
            0.24,
            _shorten(metric_value, 28),
            size=_font_size_for_length(
                metric_value.replace("\n", " "),
                6.2,
                [(10, 9), (16, 8), (24, 7)],
            ),
            bold=True,
            color=BLACK,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0.02,
        )

    _add_text(
        slide,
        0.22,
        3.25,
        2.1,
        0.15,
        "Strategic Objective Covered",
        size=7,
        bold=True,
        color=BLACK,
    )
    objectives = _text_list(data.get("strategic_objectives", []), "Strategic objective")
    for idx in range(4):
        text = _shorten(objectives[idx], 30) if idx < len(objectives) else ""
        box = _add_box(
            slide,
            0.23,
            3.47 + idx * 0.28,
            2.14,
            0.22,
            OBJECTIVE_COLORS[idx],
            OBJECTIVE_COLORS[idx],
        )
        if text:
            _add_text(
                slide,
                0.3,
                3.51 + idx * 0.28,
                1.95,
                0.12,
                text,
                size=7.1,
                bold=True,
                color=WHITE if idx in (0, 1) else BLACK,
                align=PP_ALIGN.CENTER,
            )

    _add_text(
        slide,
        0.22,
        4.65,
        2.15,
        0.15,
        "Scenario-based maturity uplift",
        size=7,
        bold=True,
        color=BLACK,
    )
    uplift = _text_list(data.get("maturity_uplift", []), "x1.0")[:3]
    for idx, value in enumerate(uplift):
        x = 0.8 + idx * 0.55
        _add_text(slide, x, 4.93, 0.28, 0.1, value, size=7, bold=True, align=PP_ALIGN.CENTER)
    slide.shapes.add_connector(
        1, Inches(0.45), Inches(5.15), Inches(2.35), Inches(5.15)
    ).line.color.rgb = _rgb(RED)
    slide.shapes.add_connector(
        1, Inches(2.2), Inches(5.07), Inches(2.35), Inches(5.15)
    ).line.color.rgb = _rgb(RED)
    slide.shapes.add_connector(
        1, Inches(2.2), Inches(5.23), Inches(2.35), Inches(5.15)
    ).line.color.rgb = _rgb(RED)
    _add_text(slide, 0.23, 5.2, 0.15, 0.1, "0", size=6)
    for idx, label in enumerate(["Y1", "Y2", "Y3"]):
        _add_text(slide, 0.95 + idx * 0.55, 5.23, 0.2, 0.1, label, size=6, align=PP_ALIGN.CENTER)

    _add_text(slide, 0.22, 5.58, 1.3, 0.15, "Domains impacted", size=7, bold=True)
    impacted = _text_list(data.get("domains_impacted", []), "Domain")
    for idx in range(2):
        top = 5.82 + idx * 0.32
        _add_box(slide, 0.23, top, 2.14, 0.25, MID_GREY, MID_GREY)
        num = _add_box(slide, 0.23, top, 0.18, 0.25, RED, RED)
        _add_text(
            slide,
            0.26,
            top + 0.05,
            0.1,
            0.08,
            str(idx + 1),
            size=7,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        if idx < len(impacted):
            _add_text(
                slide,
                0.45,
                top + 0.05,
                1.8,
                0.1,
                _shorten(impacted[idx], 34),
                size=7.5,
                color=WHITE,
            )


def _draw_center_panel(slide, data: dict) -> None:
    _add_text(slide, 2.9, 0.38, 1.8, 0.16, "Initiative Overview", size=9, bold=True)
    _add_text(slide, 2.95, 0.58, 2.0, 0.14, "Initiative Overview Description", size=6, italic=True, color=TEXT_GREY)
    _add_box(slide, 2.88, 0.76, 3.42, 0.94, WHITE, WHITE)
    _add_text(
        slide,
        2.95,
        0.8,
        3.2,
        0.84,
        _shorten(data.get("initiative_overview", "No overview provided."), 280),
        size=8.5,
        color=TEXT_GREY,
    )

    _add_text(slide, 2.9, 1.82, 1.8, 0.16, "Initiative Owner(s)", size=9, bold=True)
    _add_text(slide, 3.0, 2.08, 1.3, 0.14, "Initiative owner name", size=6, italic=True, color=TEXT_GREY)
    _add_text(
        slide,
        3.0,
        2.22,
        2.6,
        0.16,
        _shorten(data.get("initiative_owner", "CISO"), 42),
        size=9,
        bold=True,
        color=RED,
    )
    _add_text(slide, 3.0, 2.42, 1.6, 0.16, "Supporting Functions", size=8, bold=True)
    _add_text(
        slide,
        3.0,
        2.58,
        3.2,
        0.38,
        ", ".join(_text_list(data.get("supporting_functions", []))),
        size=7.2,
        color=TEXT_GREY,
    )

    _add_text(slide, 2.9, 3.02, 1.8, 0.16, "Initiative Projects", size=9, bold=True)
    projects = data.get("initiative_projects", []) or []
    current_top = 3.26
    for idx, project in enumerate(projects[:3], start=1):
        name = _shorten(project.get("name", f"Project {idx}"), 42)
        description = _shorten(project.get("description", ""), 120)
        _add_text(
            slide,
            2.95,
            current_top,
            2.8,
            0.14,
            f"{idx}.{name}",
            size=8,
            bold=True,
            color=RED,
        )
        _add_text(
            slide,
            3.0,
            current_top + 0.16,
            3.1,
            0.4,
            description or "Project description",
            size=7.3,
            color=TEXT_GREY,
        )
        current_top += 0.6


def _draw_right_panel(slide, data: dict) -> None:
    _add_text(slide, 6.72, 0.38, 1.8, 0.16, "Initiative Timeline", size=9, bold=True)
    _draw_timeline(slide, data.get("timeline_projects", []))

    _add_text(slide, 9.55, 3.02, 1.1, 0.16, "AI use case", size=8, bold=True)
    _add_text(
        slide,
        9.56,
        3.24,
        2.9,
        0.46,
        _shorten(
            data.get("ai_use_case", "Describe how AI can enhance delivery of this initiative."),
            180,
        ),
        size=7.3,
        color=TEXT_GREY,
    )

    _add_text(slide, 9.55, 4.25, 0.8, 0.16, "Priority", size=8, bold=True)
    _draw_priority(slide, data.get("priority", "Medium"))

    _add_text(slide, 9.55, 4.88, 1.0, 0.16, "Dependencies", size=8, bold=True)
    _add_bullets(
        slide,
        _text_list(data.get("dependencies", []), "No dependency provided"),
        9.55,
        5.08,
        3.0,
        0.62,
        size=7.2,
    )

    _add_text(slide, 9.55, 5.95, 1.4, 0.16, "Challenges addressed", size=8, bold=True)
    challenges = _text_list(data.get("challenges", []), "No challenge provided")
    for idx in range(2):
        top = 6.21 + idx * 0.31
        _add_box(slide, 9.55, top, 2.95, 0.23, MID_GREY, MID_GREY)
        _add_box(slide, 9.55, top, 0.18, 0.23, RED, RED)
        _add_text(
            slide,
            9.6,
            top + 0.04,
            0.08,
            0.09,
            str(idx + 1),
            size=6.8,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        if idx < len(challenges):
            _add_text(
                slide,
                9.78,
                top + 0.04,
                2.55,
                0.1,
                _shorten(challenges[idx], 48),
                size=7,
                color=WHITE,
            )


def _draw_timeline(slide, timeline_projects: list[dict]) -> None:
    left = 6.62
    top = 0.76
    width = 6.15
    height = 1.98
    _add_box(slide, left, top, width, height, WHITE, WHITE)

    year_width = 1.57
    quarter_width = 0.39
    _add_text(slide, 8.4, 0.72, 0.7, 0.12, "Year 1", size=7, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, 9.98, 0.72, 0.7, 0.12, "Year 2", size=7, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, 11.56, 0.72, 0.7, 0.12, "Year 3", size=7, bold=True, align=PP_ALIGN.CENTER)

    for group_idx in range(3):
        group_left = 7.94 + group_idx * year_width
        for quarter_idx, label in enumerate(["Q1", "Q2", "Q3", "Q4"]):
            cell_left = group_left + quarter_idx * quarter_width
            _add_box(slide, cell_left, 0.9, quarter_width - 0.01, 0.22, RED, RED)
            _add_text(
                slide,
                cell_left,
                0.945,
                quarter_width - 0.01,
                0.08,
                label,
                size=6.3,
                bold=True,
                color=WHITE,
                align=PP_ALIGN.CENTER,
            )

    rows = timeline_projects[:3] or [{"name": "Project", "start_quarter": 1, "end_quarter": 2}]
    row_height = 0.47
    start_y = 1.18
    for idx, project in enumerate(rows):
        row_top = start_y + idx * row_height
        _add_box(slide, 6.62, row_top, 6.15, 0.4, LIGHT_GREY, WHITE)
        _add_text(
            slide,
            6.73,
            row_top + 0.08,
            1.2,
            0.18,
            _shorten(project.get("name", f"Project {idx + 1}"), 25),
            size=7.4,
            bold=True,
            italic=True,
        )
        start_quarter = max(1, min(12, int(project.get("start_quarter", 1))))
        end_quarter = max(start_quarter, min(12, int(project.get("end_quarter", start_quarter))))
        bar_left = 7.94 + (start_quarter - 1) * quarter_width
        bar_width = max(quarter_width * (end_quarter - start_quarter + 1) - 0.03, 0.18)
        _add_box(slide, bar_left, row_top + 0.14, bar_width, 0.08, MINT, MINT)


def _draw_priority(slide, priority: str) -> None:
    labels = ["High", "Medium", "Low"]
    fills = {
        "High": DARK if priority == "High" else MID_GREY,
        "Medium": DARK if priority == "Medium" else MID_GREY,
        "Low": DARK if priority == "Low" else MID_GREY,
    }
    for idx, label in enumerate(labels):
        left = 9.55 + idx * 0.78
        _add_box(slide, left, 4.48, 0.68, 0.22, fills[label], fills[label], rounded=True)
        _add_text(
            slide,
            left,
            4.535,
            0.68,
            0.08,
            label,
            size=6.5,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
